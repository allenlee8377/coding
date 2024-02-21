from docx import Document
import os
import openpyxl

def convert_word_to_excel(word_path, excel_path):
    # 打开 Word 文档
    doc = Document(word_path)

    # 创建 Excel 文件
    wb = openpyxl.Workbook()
    ws = wb.active

    for table in doc.tables:
        for row in table.rows:
            row_data = [cell.text for cell in row.cells]
            ws.append(row_data)

    # 保存 Excel 文件
    wb.save(excel_path)
    print(f"成功将Word文档转换为Excel文件: {excel_path}")

# 输入 Word 和 Excel 文件的路径
docx_directory = input("请输入Word文档所在文件夹路径：")
docx_filename = input("请输入Word文档文件名：")
word_file_path = os.path.join(docx_directory, f'{docx_filename}'+".docx")
print(word_file_path)

excel_file_path = input("请输入要保存的Excel文件路径：")

# 调用转换函数
convert_word_to_excel(word_file_path, excel_file_path)
